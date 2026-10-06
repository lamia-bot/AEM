from __future__ import annotations

import csv
import json
import hashlib
import os
import time
from pathlib import Path
from typing import Callable

import numpy as np
from .cache import CacheManager
from .state import RecordingInfo

PREVIEW_POINTS = 5000
RANGE_PREVIEW_POINTS = 30_000
CHUNK_ROWS = 12000
CSV_PREVIEW_CHUNKS = 160
CSV_PREVIEW_CHUNK_BYTES = 128 * 1024
CSV_SCAN_BLOCK_BYTES = 64 * 1024 * 1024


def _fast_columns_dir(recording: RecordingInfo) -> Path:
    if recording.cache_dir is None:
        raise ValueError("No recording cache directory is available.")
    directory = Path(recording.cache_dir) / "fast_columns"
    directory.mkdir(parents=True, exist_ok=True)
    return directory


def _fast_time_path(recording: RecordingInfo) -> Path:
    return _fast_columns_dir(recording) / "time.bin"


def _fast_electrode_path(recording: RecordingInfo, electrode: str) -> Path:
    digest = hashlib.sha1(str(electrode).encode("utf-8")).hexdigest()[:16]
    return _fast_columns_dir(recording) / f"electrode_{digest}.bin"


def _valid_fast_file(path: Path, row_count: int) -> bool:
    try:
        return path.exists() and path.stat().st_size == int(row_count) * 8
    except OSError:
        return False


def fast_electrode_cache_ready(recording: RecordingInfo, electrode: str) -> bool:
    """Return True when memory-mappable Time + selected-electrode arrays exist."""
    if (
        recording.row_count <= 0
        or not recording.row_count_exact
        or electrode not in recording.electrode_columns
    ):
        return False
    return (
        _valid_fast_file(_fast_time_path(recording), recording.row_count)
        and _valid_fast_file(_fast_electrode_path(recording, electrode), recording.row_count)
    )


def _range_preview_dir(recording: RecordingInfo) -> Path:
    if recording.cache_dir is None:
        raise ValueError("No recording cache directory is available.")
    directory = Path(recording.cache_dir) / "range_previews"
    directory.mkdir(parents=True, exist_ok=True)
    return directory


def _range_preview_path(recording: RecordingInfo, electrode: str) -> Path:
    digest = hashlib.sha1(str(electrode).encode("utf-8")).hexdigest()[:16]
    return _range_preview_dir(recording) / f"range_{digest}.npz"


def load_range_preview(recording: RecordingInfo, electrode: str):
    """Load the enhanced ~30k-point Range preview when available.

    Returns ``(time, voltage, byte_offsets)`` or ``None``. Byte offsets are stored
    for lazy CSV recordings and let Apply Range jump close to the requested
    interval instead of scanning from the beginning of a multi-GB file.
    """
    path = _range_preview_path(recording, electrode)
    if not path.exists():
        return None
    try:
        with np.load(path, allow_pickle=False) as data:
            time_values = np.asarray(data["time"], dtype=np.float64)
            voltage_values = np.asarray(data["voltage"], dtype=np.float64)
            offsets = np.asarray(data["offsets"], dtype=np.int64)
        if time_values.size < 2 or voltage_values.size != time_values.size or offsets.size != time_values.size:
            return None
        return time_values, voltage_values, offsets
    except Exception:
        return None


def _csv_delimiter_for_recording(recording: RecordingInfo) -> str:
    delimiter = recording.csv_delimiter
    if delimiter:
        return "\t" if delimiter == "tab" else delimiter
    if recording.cache_dir is not None:
        try:
            meta = json.loads((Path(recording.cache_dir) / "metadata.json").read_text(encoding="utf-8"))
            delimiter = meta.get("delimiter", ",")
            return "\t" if delimiter == "tab" else str(delimiter)
        except Exception:
            pass
    return ","


def _persist_lazy_csv_scan_metadata(
    recording: RecordingInfo,
    electrode: str,
    *,
    row_count: int,
    finite_count: int,
    time_min: float | None,
    time_max: float | None,
) -> None:
    """Persist facts learned by the first full selected-electrode CSV scan."""
    if recording.cache_dir is None:
        return
    meta_path = Path(recording.cache_dir) / "metadata.json"
    try:
        meta = json.loads(meta_path.read_text(encoding="utf-8"))
        meta["row_count"] = int(row_count)
        meta["row_count_exact"] = True
        meta["last_used"] = time.time()
        if time_min is not None:
            meta["time_min"] = float(time_min)
        if time_max is not None:
            meta["time_max"] = float(time_max)
        has_data = meta.get("electrode_has_data")
        if not isinstance(has_data, dict):
            has_data = {}
        has_data[str(electrode)] = bool(finite_count > 0)
        meta["electrode_has_data"] = has_data
        tmp = meta_path.with_suffix(".json.tmp")
        tmp.write_text(json.dumps(meta, indent=2), encoding="utf-8")
        os.replace(tmp, meta_path)
    except Exception:
        # Sidecars are already valid even if metadata refreshing fails. Reopening
        # will simply rebuild/validate the small cache next time.
        pass


def _prepare_fast_csv_electrode_cache(
    recording: RecordingInfo,
    electrode: str,
    progress: Callable[[int], None] | None = None,
) -> None:
    """Scan a large CSV once and extract only Time + the selected electrode.

    Build 19 intentionally avoids creating a 64-column Parquet copy of a multi-GB
    CSV. The initial open only makes a sparse preview. Once the user chooses an
    electrode, Arrow streams the source exactly once and writes two raw float64
    sidecars (Time and that electrode). All later Range operations are memory-map
    slices, preserving the fast Build 14/15 behavior downstream.
    """
    if recording.source_path is None or recording.cache_dir is None:
        raise ValueError("No CSV source/cache is available.")
    if electrode not in recording.electrode_columns:
        raise ValueError(f"Unknown electrode: {electrode}")

    try:
        import pyarrow as pa
        import pyarrow.csv as pacsv
    except ImportError as exc:
        raise RuntimeError(
            "pyarrow is required for large CSV electrode preparation. Install requirements.txt first."
        ) from exc

    source_path = Path(recording.source_path)
    delimiter = _csv_delimiter_for_recording(recording)
    electrode_index = recording.electrode_columns.index(electrode) + 1
    electrode_input_name = f"f{electrode_index}"
    time_input_name = "f0"

    time_path = _fast_time_path(recording)
    electrode_path = _fast_electrode_path(recording, electrode)
    time_is_valid = (
        recording.row_count_exact
        and recording.row_count > 0
        and _valid_fast_file(time_path, recording.row_count)
    )
    electrode_is_valid = (
        recording.row_count_exact
        and recording.row_count > 0
        and _valid_fast_file(electrode_path, recording.row_count)
    )
    if time_is_valid and electrode_is_valid:
        if progress:
            progress(100)
        return

    need_time = not time_is_valid
    need_electrode = not electrode_is_valid
    tmp_time = time_path.with_name(time_path.name + ".tmp")
    tmp_electrode = electrode_path.with_name(electrode_path.name + ".tmp")
    for tmp in (tmp_time, tmp_electrode):
        try:
            tmp.unlink()
        except FileNotFoundError:
            pass

    include_columns: list[str] = []
    column_types: dict[str, object] = {}
    if need_time:
        include_columns.append(time_input_name)
        column_types[time_input_name] = pa.float64()
    if need_electrode:
        include_columns.append(electrode_input_name)
        column_types[electrode_input_name] = pa.float64()

    source_size = max(source_path.stat().st_size, 1)
    stream = None
    time_handle = None
    electrode_handle = None
    total = 0
    finite_count = 0
    time_min = None
    time_max = None
    last_pct = -1
    try:
        stream = pa.input_stream(str(source_path))
        read_options = pacsv.ReadOptions(
            use_threads=True,
            block_size=CSV_SCAN_BLOCK_BYTES,
            skip_rows=1,
            autogenerate_column_names=True,
        )
        parse_options = pacsv.ParseOptions(delimiter=delimiter)
        convert_options = pacsv.ConvertOptions(
            include_columns=include_columns,
            column_types=column_types,
            strings_can_be_null=True,
            include_missing_columns=True,
            null_values=["", "NA", "N/A", "NaN", "nan", "NULL", "null"],
        )
        reader = pacsv.open_csv(
            stream,
            read_options=read_options,
            parse_options=parse_options,
            convert_options=convert_options,
        )
        if need_time:
            time_handle = open(tmp_time, "wb")
        if need_electrode:
            electrode_handle = open(tmp_electrode, "wb")

        for batch in reader:
            n = int(batch.num_rows)
            if n <= 0:
                continue
            if need_time:
                idx = batch.schema.get_field_index(time_input_name)
                time_values = np.asarray(
                    batch.column(idx).to_numpy(zero_copy_only=False), dtype=np.float64
                )
                time_values.tofile(time_handle)
                finite_time = time_values[np.isfinite(time_values)]
                if finite_time.size:
                    batch_min = float(np.min(finite_time))
                    batch_max = float(np.max(finite_time))
                    time_min = batch_min if time_min is None else min(time_min, batch_min)
                    time_max = batch_max if time_max is None else max(time_max, batch_max)
            if need_electrode:
                idx = batch.schema.get_field_index(electrode_input_name)
                electrode_values = np.asarray(
                    batch.column(idx).to_numpy(zero_copy_only=False), dtype=np.float64
                )
                finite_count += int(np.isfinite(electrode_values).sum())
                electrode_values.tofile(electrode_handle)
            total += n

            if progress:
                try:
                    consumed = int(stream.tell())
                    pct = min(99, max(1, int(100 * min(consumed, source_size) / source_size)))
                except Exception:
                    # Still provide visible movement on Arrow builds whose stream
                    # wrapper does not expose byte position.
                    pct = min(99, 1 + int(7 * np.log10(max(total, 1))))
                if pct != last_pct:
                    progress(pct)
                    last_pct = pct

        if total <= 0:
            raise ValueError("The selected CSV contains headers but no data rows.")
        if recording.row_count_exact and recording.row_count > 0 and total != recording.row_count:
            raise RuntimeError(
                f"CSV row count changed while preparing the electrode ({total:,} vs {recording.row_count:,})."
            )
    except Exception:
        for tmp in (tmp_time, tmp_electrode):
            try:
                tmp.unlink()
            except FileNotFoundError:
                pass
        raise
    finally:
        if time_handle is not None:
            time_handle.close()
        if electrode_handle is not None:
            electrode_handle.close()
        if stream is not None:
            try:
                stream.close()
            except Exception:
                pass

    expected_bytes = int(total) * 8
    try:
        if need_time:
            if not tmp_time.exists() or tmp_time.stat().st_size != expected_bytes:
                raise RuntimeError("Fast Time cache size is invalid.")
            os.replace(tmp_time, time_path)
        if need_electrode:
            if not tmp_electrode.exists() or tmp_electrode.stat().st_size != expected_bytes:
                raise RuntimeError("Fast electrode cache size is invalid.")
            os.replace(tmp_electrode, electrode_path)

        recording.row_count = int(total)
        recording.row_count_exact = True
        if time_min is not None:
            recording.time_min = float(time_min)
        if time_max is not None:
            recording.time_max = float(time_max)
        if need_electrode:
            recording.electrode_has_data[electrode] = bool(finite_count > 0)
        _persist_lazy_csv_scan_metadata(
            recording,
            electrode,
            row_count=total,
            finite_count=finite_count,
            time_min=time_min,
            time_max=time_max,
        )
        if progress:
            progress(100)
    except Exception:
        for tmp in (tmp_time, tmp_electrode):
            try:
                tmp.unlink()
            except FileNotFoundError:
                pass
        raise


def prepare_fast_electrode_cache(
    recording: RecordingInfo,
    electrode: str,
    progress: Callable[[int], None] | None = None,
) -> None:
    """Prepare memory-mapped Time + selected-electrode arrays.

    XLSX and older caches use their existing Parquet source. Build 19 lazy CSV
    caches scan the original CSV directly and extract only the chosen electrode.
    """
    if electrode not in recording.electrode_columns:
        raise ValueError(f"Unknown electrode: {electrode}")

    if recording.source_format.casefold() == "csv" and recording.cache_parquet is None:
        _prepare_fast_csv_electrode_cache(recording, electrode, progress)
        return

    if recording.cache_parquet is None or recording.time_column is None:
        raise ValueError("No cached recording is available.")
    if recording.row_count <= 0:
        raise ValueError("The recording has no data rows.")

    time_path = _fast_time_path(recording)
    electrode_path = _fast_electrode_path(recording, electrode)
    need_time = not _valid_fast_file(time_path, recording.row_count)
    need_electrode = not _valid_fast_file(electrode_path, recording.row_count)
    if not need_time and not need_electrode:
        if progress:
            progress(100)
        return

    try:
        import pyarrow.parquet as pq
    except ImportError as exc:
        raise RuntimeError("pyarrow is required. Install requirements.txt first.") from exc

    tmp_time = time_path.with_name(time_path.name + ".tmp")
    tmp_electrode = electrode_path.with_name(electrode_path.name + ".tmp")
    for tmp in (tmp_time, tmp_electrode):
        try:
            tmp.unlink()
        except FileNotFoundError:
            pass

    columns = []
    if need_time:
        columns.append(recording.time_column)
    if need_electrode:
        columns.append(electrode)

    parquet = pq.ParquetFile(recording.cache_parquet, memory_map=True)
    time_handle = open(tmp_time, "wb") if need_time else None
    electrode_handle = open(tmp_electrode, "wb") if need_electrode else None
    total = 0
    last_pct = -1
    try:
        for batch in parquet.iter_batches(
            batch_size=524_288,
            columns=columns,
            use_threads=True,
        ):
            n = int(batch.num_rows)
            if need_time:
                idx = batch.schema.get_field_index(recording.time_column)
                values = np.asarray(batch.column(idx).to_numpy(zero_copy_only=False), dtype=np.float64)
                values.tofile(time_handle)
            if need_electrode:
                idx = batch.schema.get_field_index(electrode)
                values = np.asarray(batch.column(idx).to_numpy(zero_copy_only=False), dtype=np.float64)
                values.tofile(electrode_handle)
            total += n
            if progress and recording.row_count:
                pct = min(99, int(100 * total / recording.row_count))
                if pct != last_pct:
                    progress(pct)
                    last_pct = pct
    finally:
        if time_handle is not None:
            time_handle.close()
        if electrode_handle is not None:
            electrode_handle.close()

    expected_bytes = int(recording.row_count) * 8
    try:
        if total != int(recording.row_count):
            raise RuntimeError(
                f"Fast electrode cache row count mismatch ({total:,} vs {recording.row_count:,})."
            )
        if need_time:
            if tmp_time.stat().st_size != expected_bytes:
                raise RuntimeError("Fast Time cache size is invalid.")
            os.replace(tmp_time, time_path)
        if need_electrode:
            if tmp_electrode.stat().st_size != expected_bytes:
                raise RuntimeError("Fast electrode cache size is invalid.")
            os.replace(tmp_electrode, electrode_path)
        if progress:
            progress(100)
    except Exception:
        for tmp in (tmp_time, tmp_electrode):
            try:
                tmp.unlink()
            except FileNotFoundError:
                pass
        raise

def _load_fast_electrode_range(
    recording: RecordingInfo, electrode: str, start: float, end: float
):
    """Return a range from memory-mapped sidecars, or None when unavailable."""
    if not fast_electrode_cache_ready(recording, electrode):
        return None
    time_values = np.memmap(
        _fast_time_path(recording), dtype=np.float64, mode="r", shape=(recording.row_count,)
    )
    voltage_values = np.memmap(
        _fast_electrode_path(recording, electrode),
        dtype=np.float64,
        mode="r",
        shape=(recording.row_count,),
    )
    # Machine recordings have an increasing Time axis. Search the mapped Time
    # vector, then copy only the requested interval into analysis memory.
    left = int(np.searchsorted(time_values, start, side="left"))
    right = int(np.searchsorted(time_values, end, side="right"))
    left = max(0, min(left, recording.row_count))
    right = max(left, min(right, recording.row_count))
    if right <= left:
        return np.array([], dtype=np.float64), np.array([], dtype=np.float64)
    return (
        np.asarray(time_values[left:right], dtype=np.float64).copy(),
        np.asarray(voltage_values[left:right], dtype=np.float64).copy(),
    )


def _safe_float(value):
    if value is None or value == "":
        return np.nan
    try:
        return float(value)
    except (TypeError, ValueError):
        return np.nan


def _unique_headers(raw_headers: list[object]) -> list[str]:
    headers: list[str] = []
    seen: dict[str, int] = {}
    for i, raw in enumerate(raw_headers):
        base = str(raw).strip() if raw not in (None, "") else f"Column_{i + 1}"
        count = seen.get(base, 0)
        seen[base] = count + 1
        headers.append(base if count == 0 else f"{base}_{count + 1}")
    return headers


def _parquet_time_bounds(parquet_path: Path, time_column: str):
    """Read min/max Time from Parquet row-group metadata without loading samples."""
    try:
        import pyarrow.parquet as pq
        parquet = pq.ParquetFile(parquet_path)
        time_idx = parquet.schema_arrow.get_field_index(time_column)
        mins = []
        maxs = []
        if time_idx >= 0:
            for i in range(parquet.num_row_groups):
                stats = parquet.metadata.row_group(i).column(time_idx).statistics
                if stats is not None and stats.has_min_max:
                    mins.append(float(stats.min))
                    maxs.append(float(stats.max))
        if mins and maxs:
            return float(min(mins)), float(max(maxs))
    except Exception:
        pass
    return None, None


def load_cached_recording(directory: Path, cache_manager: CacheManager) -> RecordingInfo:
    meta = json.loads((directory / "metadata.json").read_text(encoding="utf-8"))
    npz = np.load(directory / "preview.npz", allow_pickle=False)
    preview_time = npz["time"]
    preview_signals = {
        column: npz[f"e_{i}"] for i, column in enumerate(meta["electrode_columns"])
    }
    electrode_has_data = meta.get("electrode_has_data")
    if not isinstance(electrode_has_data, dict):
        electrode_has_data = {
            column: bool(np.isfinite(np.asarray(preview_signals.get(column, []), dtype=float)).any())
            for column in meta["electrode_columns"]
        }
    else:
        electrode_has_data = {str(k): bool(v) for k, v in electrode_has_data.items()}

    parquet_candidate = directory / "recording.parquet"
    parquet_path = parquet_candidate if parquet_candidate.exists() else None
    time_min = meta.get("time_min")
    time_max = meta.get("time_max")
    if (time_min is None or time_max is None) and parquet_path is not None:
        time_min, time_max = _parquet_time_bounds(parquet_path, meta["time_column"])
    if time_min is None or time_max is None:
        finite_preview = preview_time[np.isfinite(preview_time)]
        if finite_preview.size:
            time_min = float(np.min(finite_preview))
            time_max = float(np.max(finite_preview))

    delimiter = meta.get("delimiter")
    if delimiter == "tab":
        delimiter = "\t"
    cache_manager.touch(directory)
    return RecordingInfo(
        source_path=Path(meta["source_path"]),
        cache_dir=directory,
        cache_parquet=parquet_path,
        headers=meta["headers"],
        time_column=meta["time_column"],
        electrode_columns=meta["electrode_columns"],
        row_count=int(meta.get("row_count", 0)),
        row_count_exact=bool(meta.get("row_count_exact", True)),
        source_size=int(meta["source_size"]),
        source_format=str(meta.get("source_format") or Path(meta["source_path"]).suffix.lstrip(".").lower()),
        csv_delimiter=None if delimiter is None else str(delimiter),
        time_min=None if time_min is None else float(time_min),
        time_max=None if time_max is None else float(time_max),
        preview_time=preview_time,
        preview_signals=preview_signals,
        electrode_has_data=electrode_has_data,
    )

def _cache_output_paths(source_path: Path, cache_manager: CacheManager):
    directory = cache_manager.entry_dir(source_path)
    directory.mkdir(parents=True, exist_ok=True)
    parquet_path = directory / "recording.parquet"
    preview_path = directory / "preview.npz"
    meta_path = directory / "metadata.json"
    fast_dir = directory / "fast_columns"
    fast_dir.mkdir(parents=True, exist_ok=True)
    fast_time_path = fast_dir / "time.bin"
    fast_time_tmp = fast_dir / "time.bin.tmp"
    try:
        fast_time_tmp.unlink()
    except FileNotFoundError:
        pass
    return directory, parquet_path, preview_path, meta_path, fast_time_path, fast_time_tmp


def _cleanup_failed_cache(*paths: Path) -> None:
    for path in paths:
        try:
            if path.exists():
                path.unlink()
        except OSError:
            pass


def _write_cache_metadata(
    source_path: Path,
    meta_path: Path,
    *,
    headers: list[str],
    time_column: str,
    electrode_columns: list[str],
    electrode_finite_counts: list[int],
    row_count: int,
    time_min: float | None,
    time_max: float | None,
    source_format: str,
    delimiter: str | None = None,
    row_count_exact: bool = True,
    cache_mode: str | None = None,
) -> None:
    stat = source_path.stat()
    now = time.time()
    metadata = {
        "source_name": source_path.name,
        "source_path": str(source_path),
        "source_size": stat.st_size,
        "source_fingerprint": {
            "resolved_path": str(source_path),
            "size": stat.st_size,
            "mtime_ns": stat.st_mtime_ns,
        },
        "created": now,
        "last_used": now,
        "source_format": source_format,
        "headers": headers,
        "time_column": time_column,
        "electrode_columns": electrode_columns,
        "electrode_has_data": {
            column: bool(electrode_finite_counts[i] > 0)
            for i, column in enumerate(electrode_columns)
        },
        "row_count": int(row_count),
        "row_count_exact": bool(row_count_exact),
        "time_min": time_min,
        "time_max": time_max,
    }
    if delimiter is not None:
        metadata["delimiter"] = "tab" if delimiter == "\t" else delimiter
    if cache_mode is not None:
        metadata["cache_mode"] = cache_mode
    meta_path.write_text(json.dumps(metadata, indent=2), encoding="utf-8")


def _save_preview(
    preview_path: Path,
    preview_time,
    preview_electrodes,
) -> None:
    """Save a capped preview without changing any full-resolution analysis data."""
    n = len(preview_time)
    if n > PREVIEW_POINTS:
        keep = np.linspace(0, n - 1, PREVIEW_POINTS, dtype=int)
        preview_time = [preview_time[i] for i in keep]
        preview_electrodes = [
            [values[i] for i in keep]
            for values in preview_electrodes
        ]
    payload = {"time": np.asarray(preview_time, dtype=np.float64)}
    for i, values in enumerate(preview_electrodes):
        payload[f"e_{i}"] = np.asarray(values, dtype=np.float64)
    np.savez_compressed(preview_path, **payload)


def _has_xlsx_zip_signature(source_path: Path) -> bool:
    """Return True when the file begins with the ZIP signature used by XLSX files."""
    try:
        with source_path.open("rb") as handle:
            return handle.read(4) == b"PK\x03\x04"
    except OSError:
        return False


def _looks_like_delimited_text(source_path: Path) -> bool:
    """Lightweight content check used only to route a recording to CSV vs XLSX."""
    try:
        with source_path.open("rb") as handle:
            raw = handle.read(131072)
    except OSError:
        return False
    if not raw or b"\x00" in raw:
        return False
    try:
        text = raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        try:
            text = raw.decode("latin-1")
        except UnicodeDecodeError:
            return False
    lines = [line for line in text.splitlines() if line.strip()]
    if not lines:
        return False
    first = lines[0]
    return max(first.count(","), first.count("\t"), first.count(";")) >= 1


def detect_recording_format(source_path: Path) -> str:
    """Return ``csv`` or ``xlsx`` without ever asking openpyxl to inspect CSV text.

    Build 18 deliberately makes the routing explicit and content-aware. Machine
    recordings are primarily CSV, but XLSX remains supported. Common text extensions
    (.csv/.tsv/.txt) always take the CSV path. For extensionless or oddly named files,
    the first bytes are inspected as a safe fallback.
    """
    source_path = Path(source_path)
    suffix = source_path.suffix.casefold().strip()

    if suffix in {".csv", ".tsv", ".txt"}:
        if _has_xlsx_zip_signature(source_path):
            raise ValueError(
                "This file is named like a CSV/text recording but its contents look like an XLSX workbook. "
                "Rename it with the .xlsx extension and select it again."
            )
        return "csv"

    if suffix == ".xlsx":
        if _has_xlsx_zip_signature(source_path):
            return "xlsx"
        if _looks_like_delimited_text(source_path):
            # Some export workflows leave an .xlsx name on plain-text data. Treat
            # the actual content as authoritative so the user is not sent to openpyxl.
            return "csv"
        return "xlsx"

    # Fallback for extensionless/oddly named exports.
    if _has_xlsx_zip_signature(source_path):
        return "xlsx"
    if _looks_like_delimited_text(source_path):
        return "csv"
    raise ValueError(
        "Unsupported recording format. Choose a CSV/delimited-text file or an .xlsx workbook."
    )


def _detect_csv_structure(source_path: Path) -> tuple[str, list[object]]:
    """Detect comma/tab/semicolon separation and return the header row.

    MEA machine files are expected to contain one header row, then numeric data.
    The first column is Time and the next up to 64 columns are electrodes. Exact
    header wording is deliberately ignored by the scientific pipeline.
    """
    with source_path.open("r", encoding="utf-8-sig", errors="replace", newline="") as handle:
        sample = handle.read(131072)
        if not sample.strip():
            raise ValueError("The selected CSV file is empty.")
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
            delimiter = dialect.delimiter
        except csv.Error:
            first_line = next((line for line in sample.splitlines() if line.strip()), "")
            counts = {candidate: first_line.count(candidate) for candidate in (",", "\t", ";")}
            delimiter = max(counts, key=counts.get)
            if counts[delimiter] == 0:
                raise ValueError("Could not determine the CSV delimiter (comma, tab, or semicolon).")

        handle.seek(0)
        reader = csv.reader(handle, delimiter=delimiter)
        try:
            raw_headers = next(reader)
        except StopIteration as exc:
            raise ValueError("The selected CSV file has no header row.") from exc

    if len(raw_headers) < 2:
        raise ValueError("The recording must contain Time plus at least one electrode column.")
    # Time + at most 64 electrodes. Any machine metadata columns to the right
    # are intentionally ignored rather than being offered as electrodes.
    return delimiter, list(raw_headers[:65])


def _parse_sampled_csv_line(raw_line: bytes, delimiter: str, n_cols: int) -> list[float] | None:
    if not raw_line.strip():
        return None
    try:
        text = raw_line.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = raw_line.decode("latin-1", errors="replace")
    try:
        row = next(csv.reader([text], delimiter=delimiter))
    except Exception:
        return None
    if not row:
        return None
    if len(row) < n_cols:
        row.extend([""] * (n_cols - len(row)))
    values = [_safe_float(row[i]) for i in range(n_cols)]
    if not np.isfinite(values[0]):
        return None
    return values


def _sample_csv_preview(
    source_path: Path,
    delimiter: str,
    n_cols: int,
    progress: Callable[[int, str], None] | None = None,
):
    """Read ~20 MB spread across the CSV instead of scanning the whole file.

    This is the key Build 19 change. A 2-3 GB recording can reach Electrode
    selection after a sparse, file-wide preview rather than a full conversion.
    """
    source_size = source_path.stat().st_size
    with source_path.open("rb") as handle:
        header = handle.readline()
        data_start = handle.tell()
        if not header or data_start >= source_size:
            raise ValueError("The selected CSV contains a header but no data rows.")

        data_bytes = max(1, source_size - data_start)
        max_preview_bytes = CSV_PREVIEW_CHUNKS * CSV_PREVIEW_CHUNK_BYTES
        if data_bytes <= max_preview_bytes:
            starts = list(range(data_start, source_size, CSV_PREVIEW_CHUNK_BYTES))
        else:
            last_start = max(data_start, source_size - CSV_PREVIEW_CHUNK_BYTES)
            starts = np.linspace(data_start, last_start, CSV_PREVIEW_CHUNKS, dtype=np.int64).tolist()
            starts = sorted(set(int(v) for v in starts))
        if not starts:
            starts = [data_start]

        per_chunk_target = max(1, int(np.ceil(PREVIEW_POINTS / len(starts))))
        sampled_rows: list[list[float]] = []
        sampled_line_lengths: list[int] = []
        density_rows = 0
        density_bytes = 0

        for chunk_index, start in enumerate(starts):
            handle.seek(int(start))
            if start > data_start:
                # Discard the partial row containing this byte position.
                handle.readline()
            blob = handle.read(CSV_PREVIEW_CHUNK_BYTES)
            if handle.tell() < source_size:
                # Complete the final row in the sampled block.
                blob += handle.readline()
            lines = [line for line in blob.splitlines(keepends=True) if line.strip()]
            if lines:
                density_rows += len(lines)
                density_bytes += sum(len(line) for line in lines)
                # Row-length density gives a useful approximate sample count for
                # the Recording/preview UI without paying for a full-file row count.
                if len(sampled_line_lengths) < 30_000:
                    step_len = max(1, len(lines) // 200)
                    sampled_line_lengths.extend(len(line) for line in lines[::step_len])

                take_n = min(per_chunk_target, len(lines))
                if take_n == 1:
                    indices = [len(lines) // 2]
                else:
                    indices = np.linspace(0, len(lines) - 1, take_n, dtype=int).tolist()
                for idx in indices:
                    values = _parse_sampled_csv_line(lines[int(idx)], delimiter, n_cols)
                    if values is not None:
                        sampled_rows.append(values)

            if progress:
                pct = 8 + int(84 * (chunk_index + 1) / len(starts))
                progress(min(92, pct), f"Sampling CSV preview… {chunk_index + 1}/{len(starts)}")

    if len(sampled_rows) < 2:
        raise ValueError("Could not read enough numeric rows from the CSV preview.")

    # File order is already chronological for MEA exports. Overlapping sample
    # blocks can duplicate a few points; cap/downsample later just as in older builds.
    preview_time = [row[0] for row in sampled_rows]
    preview_electrodes = [
        [row[col] for row in sampled_rows]
        for col in range(1, n_cols)
    ]
    electrode_finite_counts = [
        int(np.isfinite(np.asarray(values, dtype=np.float64)).sum())
        for values in preview_electrodes
    ]
    finite_time = np.asarray(preview_time, dtype=np.float64)
    finite_time = finite_time[np.isfinite(finite_time)]
    if finite_time.size < 2:
        raise ValueError("The CSV preview has no usable Time axis.")

    if density_rows > 0 and density_bytes > 0:
        rows_per_byte = float(density_rows) / float(density_bytes)
        row_count_estimate = max(
            len(sampled_rows),
            int(round((source_size - data_start) * rows_per_byte)),
        )
    elif sampled_line_lengths:
        median_row_bytes = max(1.0, float(np.median(np.asarray(sampled_line_lengths, dtype=float))))
        row_count_estimate = max(len(sampled_rows), int(round((source_size - data_start) / median_row_bytes)))
    else:
        row_count_estimate = len(sampled_rows)

    return (
        preview_time,
        preview_electrodes,
        electrode_finite_counts,
        int(row_count_estimate),
        float(np.min(finite_time)),
        float(np.max(finite_time)),
    )


def _parse_csv_time_and_electrode(raw_line: bytes, delimiter: str, electrode_index: int):
    """Fast numeric-row parser for one Time + one electrode value.

    MEA data rows are ordinary delimited numeric text. The common path uses a
    byte split (implemented in C) and falls back to csv.reader only if quoting is
    present. ``electrode_index`` is the zero-based column index in the source.
    """
    if not raw_line.strip():
        return None
    try:
        if b'"' not in raw_line:
            parts = raw_line.rstrip(b"\r\n").split(delimiter.encode("utf-8"), electrode_index + 1)
            if len(parts) <= electrode_index:
                return None
            time_value = float(parts[0]) if parts[0] else np.nan
            electrode_value = float(parts[electrode_index]) if parts[electrode_index] else np.nan
        else:
            try:
                text = raw_line.decode("utf-8-sig")
            except UnicodeDecodeError:
                text = raw_line.decode("latin-1", errors="replace")
            row = next(csv.reader([text], delimiter=delimiter))
            if len(row) <= electrode_index:
                return None
            time_value = _safe_float(row[0])
            electrode_value = _safe_float(row[electrode_index])
    except (ValueError, TypeError, csv.Error):
        return None
    if not np.isfinite(time_value):
        return None
    return float(time_value), float(electrode_value)


def prepare_range_preview(
    recording: RecordingInfo,
    electrode: str,
    progress: Callable[[int], None] | None = None,
):
    """Create/load a ~30k-point preview for the Range page.

    For lazy CSV recordings this reads only the same sparse ~20 MB spread across
    the source file used by the fast initial preview, but parses only Time + the
    selected electrode. It never builds a full-resolution electrode cache. Exact
    source-byte offsets are saved alongside the preview so later full-resolution
    extraction can jump close to the selected start time.
    """
    if electrode not in recording.electrode_columns:
        raise ValueError(f"Unknown electrode: {electrode}")

    cached = load_range_preview(recording, electrode)
    if cached is not None:
        if progress:
            progress(100)
        return cached

    # XLSX / Parquet recordings keep their established cache path. Their existing
    # preview is already cheap to display; no new scientific or I/O behavior is
    # required for Build 20.
    if recording.source_format.casefold() != "csv" or recording.source_path is None:
        time_values = np.asarray(recording.preview_time, dtype=np.float64)
        voltage_values = np.asarray(recording.preview_signals.get(electrode, []), dtype=np.float64)
        offsets = np.full(time_values.shape, -1, dtype=np.int64)
        return time_values, voltage_values, offsets

    source_path = Path(recording.source_path)
    delimiter = _csv_delimiter_for_recording(recording)
    electrode_index = recording.electrode_columns.index(electrode) + 1
    source_size = source_path.stat().st_size

    sampled: list[tuple[int, float, float]] = []
    with source_path.open("rb") as handle:
        header = handle.readline()
        data_start = handle.tell()
        if not header or data_start >= source_size:
            raise ValueError("The selected CSV contains a header but no data rows.")

        data_bytes = max(1, source_size - data_start)
        max_preview_bytes = CSV_PREVIEW_CHUNKS * CSV_PREVIEW_CHUNK_BYTES
        if data_bytes <= max_preview_bytes:
            starts = list(range(data_start, source_size, CSV_PREVIEW_CHUNK_BYTES))
        else:
            last_start = max(data_start, source_size - CSV_PREVIEW_CHUNK_BYTES)
            starts = np.linspace(data_start, last_start, CSV_PREVIEW_CHUNKS, dtype=np.int64).tolist()
            starts = sorted(set(int(v) for v in starts))
        if not starts:
            starts = [data_start]

        per_chunk_target = max(1, int(np.ceil(RANGE_PREVIEW_POINTS / len(starts))))
        for chunk_index, start in enumerate(starts):
            handle.seek(int(start))
            if start > data_start:
                handle.readline()  # discard partial row
            block_start = handle.tell()
            blob = handle.read(CSV_PREVIEW_CHUNK_BYTES)
            if handle.tell() < source_size:
                blob += handle.readline()  # complete final row

            raw_lines = []
            offsets = []
            cursor = 0
            for line in blob.splitlines(keepends=True):
                line_offset = block_start + cursor
                cursor += len(line)
                if line.strip():
                    raw_lines.append(line)
                    offsets.append(line_offset)
            if raw_lines:
                take_n = min(per_chunk_target, len(raw_lines))
                if take_n == 1:
                    indices = [len(raw_lines) // 2]
                else:
                    indices = np.linspace(0, len(raw_lines) - 1, take_n, dtype=int).tolist()

                for idx in indices:
                    i = int(idx)
                    parsed = _parse_csv_time_and_electrode(raw_lines[i], delimiter, electrode_index)
                    if parsed is None:
                        continue
                    time_value, voltage_value = parsed
                    sampled.append((int(offsets[i]), time_value, voltage_value))

            if progress:
                progress(min(99, int(100 * (chunk_index + 1) / len(starts))))

    if len(sampled) < 2:
        raise ValueError("Could not create a usable Range preview for the selected electrode.")

    # Overlapping sparse blocks can capture the same row twice. Source-byte
    # position is the stable identity and preserves chronological file order.
    sampled.sort(key=lambda item: item[0])
    unique: list[tuple[int, float, float]] = []
    last_offset = None
    for item in sampled:
        if item[0] == last_offset:
            continue
        unique.append(item)
        last_offset = item[0]

    if len(unique) > RANGE_PREVIEW_POINTS:
        keep = np.linspace(0, len(unique) - 1, RANGE_PREVIEW_POINTS, dtype=int)
        unique = [unique[int(i)] for i in keep]

    byte_offsets = np.asarray([item[0] for item in unique], dtype=np.int64)
    time_values = np.asarray([item[1] for item in unique], dtype=np.float64)
    voltage_values = np.asarray([item[2] for item in unique], dtype=np.float64)

    path = _range_preview_path(recording, electrode)
    tmp = path.with_suffix(".npz.tmp")
    try:
        with open(tmp, "wb") as handle:
            np.savez(handle, time=time_values, voltage=voltage_values, offsets=byte_offsets)
        os.replace(tmp, path)
    finally:
        try:
            tmp.unlink()
        except FileNotFoundError:
            pass

    if progress:
        progress(100)
    return time_values, voltage_values, byte_offsets


def _lazy_csv_start_offset(recording: RecordingInfo, electrode: str, start: float) -> int:
    """Choose a safe row boundary shortly before ``start`` using the Range index."""
    if recording.source_path is None:
        return 0
    with Path(recording.source_path).open("rb") as handle:
        handle.readline()
        data_start = int(handle.tell())

    preview = load_range_preview(recording, electrode)
    if preview is None:
        return data_start
    times, _values, offsets = preview
    finite = np.isfinite(times) & (offsets >= data_start)
    if finite.sum() < 2:
        return data_start
    t = times[finite]
    o = offsets[finite]
    # The recordings are time ordered. If a malformed export violates that,
    # fall back to the beginning rather than risk skipping valid selected rows.
    if np.any(np.diff(t) < 0):
        return data_start
    idx = int(np.searchsorted(t, float(start), side="right") - 1)
    if idx < 0:
        return data_start
    # Start one indexed sample earlier for a generous safety margin.
    idx = max(0, idx - 1)
    return max(data_start, int(o[idx]))


def _load_lazy_csv_selected_range(
    recording: RecordingInfo,
    electrode: str,
    start: float,
    end: float,
    progress: Callable[[int], None] | None = None,
):
    """Read only the exact selected interval from a lazy multi-GB CSV.

    The Range preview's byte index lets this seek near the requested start time.
    From there rows are scanned sequentially and the inclusive v7 mask
    ``start <= time <= end`` is preserved exactly.
    """
    if recording.source_path is None:
        raise ValueError("No CSV source is available.")
    source_path = Path(recording.source_path)
    delimiter = _csv_delimiter_for_recording(recording)
    electrode_index = recording.electrode_columns.index(electrode) + 1
    seek_offset = _lazy_csv_start_offset(recording, electrode, start)

    times: list[float] = []
    values: list[float] = []
    seen_interval = False
    rows_scanned = 0
    last_pct = -1
    span = max(float(end) - float(start), 1e-12)

    with source_path.open("rb", buffering=1024 * 1024) as handle:
        handle.seek(seek_offset)
        # Offsets stored by prepare_range_preview always point to a row start.
        # For a missing/stale preview the fallback is directly after the header.
        while True:
            raw_line = handle.readline()
            if not raw_line:
                break
            parsed = _parse_csv_time_and_electrode(raw_line, delimiter, electrode_index)
            if parsed is None:
                continue
            time_value, voltage_value = parsed
            rows_scanned += 1

            if time_value < start:
                continue
            if time_value > end:
                if seen_interval:
                    break
                # If the sparse index somehow started too late, restart from the
                # data beginning once rather than silently returning an empty range.
                if seek_offset > 0 and not times:
                    with source_path.open("rb") as restart:
                        restart.readline()
                        data_start = restart.tell()
                    if seek_offset > data_start:
                        return _load_lazy_csv_selected_range_from_offset(
                            recording, electrode, start, end, data_start, progress
                        )
                break

            seen_interval = True
            times.append(time_value)
            values.append(voltage_value)

            if progress and rows_scanned % 20_000 == 0:
                pct = min(99, max(1, int(100 * (time_value - start) / span)))
                if pct != last_pct:
                    progress(pct)
                    last_pct = pct

    if progress:
        progress(100)
    return np.asarray(times, dtype=np.float64), np.asarray(values, dtype=np.float64)


def _load_lazy_csv_selected_range_from_offset(
    recording: RecordingInfo,
    electrode: str,
    start: float,
    end: float,
    offset: int,
    progress: Callable[[int], None] | None = None,
):
    """Conservative fallback used only if a sparse byte index appears stale."""
    source_path = Path(recording.source_path)
    delimiter = _csv_delimiter_for_recording(recording)
    electrode_index = recording.electrode_columns.index(electrode) + 1
    times: list[float] = []
    values: list[float] = []
    span = max(float(end) - float(start), 1e-12)
    last_pct = -1
    rows_scanned = 0
    with source_path.open("rb", buffering=1024 * 1024) as handle:
        handle.seek(int(offset))
        for raw_line in handle:
            parsed = _parse_csv_time_and_electrode(raw_line, delimiter, electrode_index)
            if parsed is None:
                continue
            time_value, voltage_value = parsed
            rows_scanned += 1
            if time_value < start:
                continue
            if time_value > end:
                break
            times.append(time_value)
            values.append(voltage_value)
            if progress and rows_scanned % 20_000 == 0:
                pct = min(99, max(1, int(100 * (time_value - start) / span)))
                if pct != last_pct:
                    progress(pct)
                    last_pct = pct
    if progress:
        progress(100)
    return np.asarray(times, dtype=np.float64), np.asarray(values, dtype=np.float64)


def prepare_csv_cache(
    source_path: Path,
    cache_manager: CacheManager,
    progress: Callable[[int, str], None] | None = None,
) -> RecordingInfo:
    """Prepare a *lazy* cache for a potentially multi-GB MEA CSV.

    Build 18 converted every CSV column to Parquet before the GUI could proceed.
    For real 2-3 GB MEA exports that duplicated gigabytes of work and made opening
    appear hung. Build 19 instead samples a file-wide preview (roughly 20 MB),
    writes only preview/metadata, and returns immediately. The selected electrode
    is extracted from the original CSV in one background scan after selection.
    """
    source_path = Path(source_path).resolve()
    if _has_xlsx_zip_signature(source_path):
        raise ValueError(
            "The selected file contains an XLSX workbook, not CSV text. Rename it with the .xlsx extension and select it again."
        )

    cached = cache_manager.valid_entry(source_path)
    if cached:
        if progress:
            progress(100, "Fast-access cache found")
        return load_cached_recording(cached, cache_manager)

    if progress:
        progress(2, "Reading CSV header")
    delimiter, raw_headers = _detect_csv_structure(source_path)
    headers = _unique_headers(raw_headers)
    time_column = headers[0]
    electrode_columns = headers[1:]
    n_cols = len(headers)

    directory, parquet_path, preview_path, meta_path, fast_time_path, fast_time_tmp = _cache_output_paths(
        source_path, cache_manager
    )
    # A cancelled/failed Build 17/18 conversion may have left a huge partial
    # Parquet file. Build 19 does not need it for a new lazy CSV cache.
    _cleanup_failed_cache(parquet_path, preview_path, meta_path, fast_time_tmp)
    if not meta_path.exists():
        try:
            for stale in (directory / "fast_columns").glob("*.tmp"):
                stale.unlink()
        except OSError:
            pass

    if progress:
        sep_name = "tab" if delimiter == "\t" else delimiter
        progress(5, f"Opening CSV quickly ({sep_name} separated)")

    (
        preview_time,
        preview_electrodes,
        electrode_finite_counts,
        row_count_estimate,
        time_min,
        time_max,
    ) = _sample_csv_preview(source_path, delimiter, n_cols, progress)

    _save_preview(preview_path, preview_time, preview_electrodes)
    _write_cache_metadata(
        source_path,
        meta_path,
        headers=headers,
        time_column=time_column,
        electrode_columns=electrode_columns,
        electrode_finite_counts=electrode_finite_counts,
        row_count=row_count_estimate,
        row_count_exact=False,
        time_min=time_min,
        time_max=time_max,
        source_format="csv",
        delimiter=delimiter,
        cache_mode="csv_lazy",
    )
    if progress:
        progress(100, "CSV preview ready — choose an electrode")
    return load_cached_recording(directory, cache_manager)

def prepare_excel_cache(
    source_path: Path,
    cache_manager: CacheManager,
    progress: Callable[[int, str], None] | None = None,
) -> RecordingInfo:
    """Stream the first XLSX sheet into Parquet and create lightweight previews.

    XLSX remains supported as a compatibility path. The machine format is assumed
    to be first column = Time, followed by up to 64 electrode columns. Header names
    are display labels only.
    """
    source_path = Path(source_path).resolve()
    if source_path.suffix.casefold() != ".xlsx" and not _has_xlsx_zip_signature(source_path):
        raise ValueError(
            "This is not an XLSX workbook. CSV recordings must be opened through the CSV loader."
        )
    if not _has_xlsx_zip_signature(source_path):
        if _looks_like_delimited_text(source_path):
            raise ValueError(
                "This file contains delimited text rather than an XLSX workbook. Please use/save it as .csv."
            )
        raise ValueError("The selected .xlsx file is not a valid XLSX workbook.")

    try:
        import pyarrow as pa
        import pyarrow.parquet as pq
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError(
            "pyarrow and openpyxl are required for XLSX recording support."
        ) from exc

    cached = cache_manager.valid_entry(source_path)
    if cached:
        if progress:
            progress(100, "Fast-access cache found")
        return load_cached_recording(cached, cache_manager)

    directory, parquet_path, preview_path, meta_path, fast_time_path, fast_time_tmp = _cache_output_paths(
        source_path, cache_manager
    )

    if progress:
        progress(2, "Opening Excel workbook")

    wb = load_workbook(source_path, read_only=True, data_only=True)
    try:
        ws = wb[wb.sheetnames[0]]
        sheet_cols = max(int(ws.max_column or 0), 2)
        n_cols = min(sheet_cols, 65)
        raw_headers = [
            cell.value
            for cell in next(ws.iter_rows(min_row=1, max_row=1, max_col=n_cols))
        ]
        if len(raw_headers) < n_cols:
            raw_headers.extend([None] * (n_cols - len(raw_headers)))
        headers = _unique_headers(raw_headers)
        time_column = headers[0]
        electrode_columns = headers[1:]
        max_row = max(int(ws.max_row or 0), 2)
        estimated_data_rows = max(max_row - 1, 1)
        preview_step = max(1, estimated_data_rows // PREVIEW_POINTS)

        preview_time: list[float] = []
        preview_electrodes: list[list[float]] = [[] for _ in electrode_columns]
        electrode_finite_counts = [0 for _ in electrode_columns]

        schema = pa.schema([(name, pa.float64()) for name in headers])
        writer = pq.ParquetWriter(
            parquet_path,
            schema,
            compression="zstd",
            use_dictionary=False,
        )
        time_fast_handle = open(fast_time_tmp, "wb")
        chunk: list[list[float]] = [[] for _ in headers]
        data_rows = 0
        time_min = None
        time_max = None
        last_percent = -1

        def flush_chunk() -> None:
            if not chunk[0]:
                return
            np.asarray(chunk[0], dtype=np.float64).tofile(time_fast_handle)
            arrays = [pa.array(values, type=pa.float64()) for values in chunk]
            table = pa.Table.from_arrays(arrays, schema=schema)
            writer.write_table(table)
            for values in chunk:
                values.clear()

        try:
            for excel_row_idx, row in enumerate(
                ws.iter_rows(min_row=2, max_col=n_cols, values_only=True), start=2
            ):
                data_rows += 1
                values = [_safe_float(row[i] if i < len(row) else None) for i in range(n_cols)]
                if np.isfinite(values[0]):
                    time_min = values[0] if time_min is None else min(time_min, values[0])
                    time_max = values[0] if time_max is None else max(time_max, values[0])
                for e_idx in range(len(electrode_columns)):
                    if np.isfinite(values[e_idx + 1]):
                        electrode_finite_counts[e_idx] += 1
                for i, value in enumerate(values):
                    chunk[i].append(value)

                if (data_rows - 1) % preview_step == 0 and len(preview_time) < PREVIEW_POINTS:
                    preview_time.append(values[0])
                    for e_idx in range(len(electrode_columns)):
                        preview_electrodes[e_idx].append(values[e_idx + 1])

                if len(chunk[0]) >= CHUNK_ROWS:
                    flush_chunk()

                if progress and data_rows % 1000 == 0:
                    pct = min(96, 5 + int(90 * data_rows / estimated_data_rows))
                    if pct != last_percent:
                        progress(pct, f"Preparing Excel recording… {data_rows:,} rows")
                        last_percent = pct
            flush_chunk()
        finally:
            writer.close()
            time_fast_handle.close()

        if data_rows == 0:
            raise ValueError("The selected Excel sheet contains headers but no data rows.")
        if fast_time_tmp.exists() and fast_time_tmp.stat().st_size == data_rows * 8:
            os.replace(fast_time_tmp, fast_time_path)
        else:
            raise RuntimeError("The Excel Time cache was not written completely.")

        _save_preview(preview_path, preview_time, preview_electrodes)
        _write_cache_metadata(
            source_path,
            meta_path,
            headers=headers,
            time_column=time_column,
            electrode_columns=electrode_columns,
            electrode_finite_counts=electrode_finite_counts,
            row_count=data_rows,
            time_min=time_min,
            time_max=time_max,
            source_format="xlsx",
        )

        if progress:
            progress(100, "Excel recording ready")
        return load_cached_recording(directory, cache_manager)
    except Exception:
        _cleanup_failed_cache(parquet_path, preview_path, meta_path, fast_time_tmp, fast_time_path)
        raise
    finally:
        wb.close()


def prepare_recording_cache(
    source_path: Path,
    cache_manager: CacheManager,
    progress: Callable[[int, str], None] | None = None,
) -> RecordingInfo:
    """Prepare CSV (preferred) or XLSX recordings for the common pipeline."""
    source_path = Path(source_path)
    source_format = detect_recording_format(source_path)
    if progress:
        progress(1, f"Detected {source_format.upper()} recording")
    if source_format == "csv":
        return prepare_csv_cache(source_path, cache_manager, progress)
    if source_format == "xlsx":
        return prepare_excel_cache(source_path, cache_manager, progress)
    raise RuntimeError(f"Unhandled recording format: {source_format}")

def load_electrode_range(
    recording: RecordingInfo,
    electrode: str,
    start: float,
    end: float,
    progress: Callable[[int], None] | None = None,
):
    """Load only the requested Time + electrode interval.

    Build 20 keeps the Range page lightweight. Lazy CSV recordings are read at
    full resolution only after Apply Range, using the sparse Range byte index to
    seek close to the chosen start time. Existing full sidecars and Parquet caches
    remain valid fast paths for older caches/XLSX recordings.
    """
    if electrode not in recording.electrode_columns:
        raise ValueError(f"Unknown electrode: {electrode}")
    start = float(start)
    end = float(end)
    if start >= end:
        raise ValueError("Start must be earlier than End.")

    # Reuse a completed Build 14-19 sidecar if one already exists. Build 20 no
    # longer creates this full-electrode cache for lazy CSV recordings.
    fast = _load_fast_electrode_range(recording, electrode, start, end)
    if fast is not None:
        if progress:
            progress(100)
        return fast

    if recording.source_format.casefold() == "csv" and recording.cache_parquet is None:
        return _load_lazy_csv_selected_range(recording, electrode, start, end, progress)

    if recording.cache_parquet is None or recording.time_column is None:
        raise ValueError("No full-resolution recording cache is available for the selected interval.")

    try:
        import pyarrow.parquet as pq
    except ImportError as exc:
        raise RuntimeError("pyarrow is required. Install requirements.txt first.") from exc

    columns = [recording.time_column, electrode]
    parquet = pq.ParquetFile(recording.cache_parquet, memory_map=True)
    time_idx = parquet.schema_arrow.get_field_index(recording.time_column)
    row_groups = []
    for row_group_index in range(parquet.num_row_groups):
        should_read = True
        if time_idx >= 0:
            stats = parquet.metadata.row_group(row_group_index).column(time_idx).statistics
            if stats is not None and stats.has_min_max:
                try:
                    rg_min = float(stats.min)
                    rg_max = float(stats.max)
                    if rg_max < start or rg_min > end:
                        should_read = False
                except (TypeError, ValueError):
                    pass
        if should_read:
            row_groups.append(row_group_index)

    if not row_groups:
        return np.array([], dtype=np.float64), np.array([], dtype=np.float64)

    # One Arrow call avoids the Python overhead of reading each row group separately.
    table = parquet.read_row_groups(row_groups, columns=columns, use_threads=True)
    time_values = np.asarray(
        table.column(recording.time_column).to_numpy(zero_copy_only=False), dtype=np.float64
    )
    voltage_values = np.asarray(
        table.column(electrode).to_numpy(zero_copy_only=False), dtype=np.float64
    )
    exact = (time_values >= start) & (time_values <= end)
    if not np.any(exact):
        return np.array([], dtype=np.float64), np.array([], dtype=np.float64)
    time_values = time_values[exact]
    voltage_values = voltage_values[exact]
    if time_values.size > 1 and np.any(np.diff(time_values) < 0):
        order = np.argsort(time_values, kind="stable")
        time_values = time_values[order]
        voltage_values = voltage_values[order]
    if progress:
        progress(100)
    return time_values, voltage_values


def load_selected_electrode(recording: RecordingInfo, electrode: str):
    """Legacy helper retained for compatibility; prefer load_electrode_range."""
    if recording.preview_time is None:
        raise ValueError("No recording preview is available.")
    finite = np.asarray(recording.preview_time, dtype=np.float64)
    finite = finite[np.isfinite(finite)]
    if finite.size < 2:
        raise ValueError("The recording has no usable time axis.")
    return load_electrode_range(recording, electrode, float(np.min(finite)), float(np.max(finite)))
